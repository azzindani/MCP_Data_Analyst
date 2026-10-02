"""An Excel export is written as a stream, so its memory does not grow with the table.

`export_data(format=excel)` on the hotel file (119,390 x 32, a 16.9 MB CSV) went through
`pd.ExcelWriter`, which keeps every cell of the sheet as an object until the file is saved: about
3 GB for 3.8 million cells. At the server's 1 GB limit the kernel killed the whole server, every
call in flight with it, and left a 0-byte workbook. The sheet is now written row by row.
"""

from __future__ import annotations

import tracemalloc
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from shared import workbook
from shared.workbook import DATA_SHEET, write_workbook


def _frame(rows: int) -> pd.DataFrame:
    rng = np.random.default_rng(0)
    return pd.DataFrame(
        {
            "n": rng.integers(0, 1000, rows),
            "amount": rng.normal(100, 10, rows).round(2),
            "grade": rng.choice(["A", "B", "C"], rows),
            "day": pd.Timestamp("2024-01-01") + pd.to_timedelta(rng.integers(0, 300, rows), unit="D"),
        }
    )


def _peak(df: pd.DataFrame, path: Path) -> int:
    tracemalloc.start()
    try:
        write_workbook(df, path)
        return tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()


class TestMemoryDoesNotGrowWithRows:
    def test_thirty_thousand_rows_do_not_hold_the_sheet(self, tmp_path):
        # Held as cell objects this frame (120,000 cells) is about 100 MB; streamed it is a chunk and
        # the per-column work, a few MB. The bound is loose on purpose: it separates the two designs.
        peak = _peak(_frame(30_000), tmp_path / "large.xlsx")
        assert peak < 15 * 1024 * 1024, f"peak was {peak:,} bytes"

    def test_every_row_is_still_there(self, tmp_path):
        out = tmp_path / "all.xlsx"
        write_workbook(_frame(12_000), out)  # more than two chunks
        sheet = openpyxl.load_workbook(out, read_only=True)[DATA_SHEET]
        assert sum(1 for _ in sheet.iter_rows()) == 12_001


class TestValuesAreWhatPandasWrote:
    @pytest.fixture
    def sheet(self, tmp_path):
        frame = pd.DataFrame(
            {
                "x": [1.5, np.nan, np.inf, -np.inf],
                "when": pd.to_datetime(["2024-03-01", None, "2024-03-03", "2024-03-04"]),
                "label": ["a", None, "c", "d"],
                "n": pd.array([1, None, 3, 4], dtype="Int64"),
            }
        )
        out = tmp_path / "v.xlsx"
        write_workbook(frame, out)
        return openpyxl.load_workbook(out)[DATA_SHEET]

    def test_a_gap_is_an_empty_cell(self, sheet):
        assert sheet["A3"].value is None and sheet["B3"].value is None
        assert sheet["C3"].value is None and sheet["D3"].value is None

    def test_infinity_is_text_because_excel_has_none(self, sheet):
        assert sheet["A4"].value == "inf" and sheet["A5"].value == "-inf"

    def test_a_date_is_a_date_with_its_format(self, sheet):
        assert sheet["B2"].value.year == 2024 and sheet["B2"].number_format == "yyyy-mm-dd"

    def test_a_number_keeps_its_format(self, sheet):
        assert sheet["A2"].value == 1.5 and sheet["A2"].number_format == "#,##0.00"
        assert sheet["D2"].value == 1 and sheet["D2"].number_format == "#,##0"

    def test_the_header_is_bold(self, sheet):
        assert sheet["A1"].font.bold is True


class TestASheetExcelCannotOpenIsRefused:
    def test_past_the_row_limit(self, tmp_path, monkeypatch):
        monkeypatch.setattr(workbook, "EXCEL_MAX_ROWS", 50)
        with pytest.raises(ValueError, match="too large for Excel"):
            write_workbook(_frame(60), tmp_path / "x.xlsx")
        monkeypatch.setattr(workbook, "EXCEL_MAX_ROWS", 70)
        write_workbook(_frame(60), tmp_path / "ok.xlsx")


class TestConvertFileStreamsToo:
    """`convert_file(output_format=excel)` held the sheet twice (`to_excel` into a BytesIO, then `getvalue()`)."""

    @pytest.fixture
    def source(self, tmp_path, monkeypatch) -> Path:
        monkeypatch.setenv("MCP_OUTPUT_DIR", str(tmp_path))
        path = tmp_path / "t.csv"
        _frame(500).assign(note=["x" if i % 7 else None for i in range(500)]).to_csv(path, index=False)
        return path

    def test_the_table_comes_back_unchanged(self, source, tmp_path):
        from servers.data_ingest.engine import convert_file

        out = tmp_path / "t.xlsx"
        result = convert_file(str(source), output_format="excel", output_path=str(out))
        assert result["success"] is True, result.get("error")
        back = pd.read_excel(out)
        original = pd.read_csv(source)
        assert back.shape == original.shape
        assert back["n"].tolist() == original["n"].tolist()
        assert back["note"].isna().sum() == original["note"].isna().sum()
        assert not [p.name for p in tmp_path.iterdir() if p.name.startswith("tmp")], "no temporary file left behind"

    def test_the_first_sheet_is_the_data(self, source, tmp_path):
        from servers.data_ingest.engine import convert_file

        out = tmp_path / "t.xlsx"
        convert_file(str(source), output_format="excel", output_path=str(out))
        assert openpyxl.load_workbook(out).sheetnames == ["Sheet1"]
